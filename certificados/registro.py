"""Pantalla del libro de constancias emitidas."""
import flet as ft
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from herramientas.registro import TIPOS, guardar_copia, listar
from ui import abrir_archivo, avisar, carpeta_descargas, desplegable

COLOR = ft.Colors.BLUE_GREY_700
COLORES_TIPO = {"Trabajo": ft.Colors.TEAL_700, "Examen": ft.Colors.ORANGE_800,
                "Licencia médica": ft.Colors.RED_700}


def vista_registro_constancias(page: ft.Page):
    buscador = ft.TextField(
        label="Buscar por docente, DNI o número (ej: 15/2026)",
        prefix_icon=ft.Icons.SEARCH, width=420, border_radius=10,
        on_change=lambda e: cargar(),
    )
    tipo = desplegable("Tipo", [("", "Todas")] + TIPOS, valor="", ancho=200)
    _elegir_original = tipo.on_select

    def al_elegir_tipo(e):
        _elegir_original(e)
        cargar()

    tipo.on_select = tipo.on_change = al_elegir_tipo

    contador = ft.Text(size=13, color=ft.Colors.GREY_600)
    vacio = ft.Text("Todavía no hay constancias registradas con ese filtro.",
                    color=ft.Colors.GREY_600, visible=False)
    tabla = ft.DataTable(
        columns=[ft.DataColumn(label=ft.Text(c, weight=ft.FontWeight.BOLD))
                 for c in ("N°", "Fecha", "Tipo", "Docente", "DNI", "Detalle", "")],
        rows=[], column_spacing=26, heading_row_height=44, data_row_min_height=46,
    )
    filas_actuales = []

    def reimprimir(id_constancia):
        try:
            ruta = guardar_copia(id_constancia, carpeta_descargas())
        except PermissionError:
            return avisar(page, "No se pudo guardar: cerrá el PDF si lo tenés abierto.")
        except Exception as error:
            print("Error al reimprimir:", error)
            return avisar(page, "No se pudo recuperar el PDF de esa constancia.")
        avisar(page, f"Copia guardada en: {ruta}")
        abrir_archivo(ruta)

    def cargar():
        filas = listar(buscador.value or "", tipo.value or None)
        filas_actuales[:] = filas
        tabla.rows = [
            ft.DataRow(cells=[
                ft.DataCell(ft.Text(f"{numero}/{anio}", weight=ft.FontWeight.BOLD)),
                ft.DataCell(ft.Text(fecha)),
                ft.DataCell(ft.Text(tipo_, color=COLORES_TIPO.get(tipo_),
                                    weight=ft.FontWeight.W_600)),
                ft.DataCell(ft.Text(docente)),
                ft.DataCell(ft.Text(dni)),
                ft.DataCell(ft.Text(detalle or "-", size=12)),
                ft.DataCell(ft.IconButton(icon=ft.Icons.PRINT, tooltip="Reimprimir (copia idéntica)",
                                          on_click=lambda e, i=id_: reimprimir(i))),
            ])
            for id_, numero, anio, fecha, tipo_, docente, dni, detalle in filas
        ]
        tabla.visible = bool(filas)
        vacio.visible = not filas
        contador.value = f"{len(filas)} constancia(s)"
        page.update()

    def exportar(e):
        if not filas_actuales:
            return avisar(page, "No hay constancias para exportar.")
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Constancias emitidas"
        hoja.append(["N°", "Fecha", "Tipo", "Docente", "DNI", "Detalle"])
        for _id, numero, anio, fecha, tipo_, docente, dni, detalle in filas_actuales:
            hoja.append([f"{numero}/{anio}", fecha, tipo_, docente, dni, detalle])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="455A64")
        for letra, ancho in zip("ABCDEF", (10, 12, 16, 32, 12, 45)):
            hoja.column_dimensions[letra].width = ancho
        ruta = carpeta_descargas() / "libro_de_constancias.xlsx"
        try:
            libro.save(ruta)
        except PermissionError:
            return avisar(page, "No se pudo guardar: cerrá el Excel si lo tenés abierto.")
        avisar(page, f"Libro exportado en: {ruta}")
        abrir_archivo(ruta)

    cargar()

    encabezado = ft.Row(spacing=15, controls=[
        ft.Icon(ft.Icons.MENU_BOOK, size=36, color=COLOR),
        ft.Column(spacing=2, controls=[
            ft.Text("Registro de constancias", size=28, weight=ft.FontWeight.BOLD),
            ft.Text("Todas las constancias emitidas, con su número. Podés reimprimir una copia "
                    "idéntica a la original.", size=14, color=ft.Colors.GREY_600),
        ]),
    ])

    return ft.Column(
        expand=True, spacing=25, scroll=ft.ScrollMode.AUTO,
        controls=[
            encabezado,
            ft.Container(
                padding=25, border_radius=12, border=ft.Border.all(1, ft.Colors.GREY_300),
                content=ft.Column(spacing=15, controls=[
                    ft.Row(wrap=True, spacing=14, vertical_alignment=ft.CrossAxisAlignment.CENTER,
                           controls=[buscador, tipo,
                                     ft.Button(content=ft.Text("Exportar a Excel"),
                                               icon=ft.Icons.TABLE_VIEW, on_click=exportar),
                                     contador]),
                    ft.Divider(),
                    ft.Row(controls=[tabla], scroll=ft.ScrollMode.AUTO),
                    vacio,
                ]),
            ),
        ],
    )
