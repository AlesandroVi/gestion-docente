"""Materias sin cubrir y planta docente por carrera (PDF y Excel)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from database.conexion import conectar
from database.consultas import FORMATO_FECHA, es_vigente, parse_fecha
from herramientas.avisos import ausentes_hoy


# ---------- DATOS ----------

def _designaciones_vigentes(hoy):
    """Designaciones vigentes hoy, con datos del docente y de la suplencia (si es suplente)."""
    conexion = conectar()
    try:
        filas = conexion.execute(
            "SELECT de.id_docente_materia, de.id_materia, de.id_docente, doc.apellido, doc.nombre, "
            "doc.dni, de.fecha_alta, de.fecha_baja, de.nro_disposicion, de.situacion, "
            "de.condicion, de.turno, de.anio, de.division, "
            "(SELECT dt.apellido || ', ' || dt.nombre FROM suplencia s "
            " JOIN designacion t ON t.id_docente_materia = s.id_designacion_titular "
            " JOIN docentes dt ON dt.id_docente = t.id_docente "
            " WHERE s.id_designacion_suplente = de.id_docente_materia) AS reemplaza "
            "FROM designacion de JOIN docentes doc ON doc.id_docente = de.id_docente"
        ).fetchall()
    finally:
        conexion.close()
    claves = ["id", "id_materia", "id_docente", "apellido", "nombre", "dni", "alta", "baja",
              "disposicion", "situacion", "condicion", "turno", "anio", "division", "reemplaza"]
    resultado = []
    for f in filas:
        fila = dict(zip(claves, f))
        alta = parse_fecha(fila["alta"])
        if es_vigente(fila["baja"], hoy) and (alta is None or alta <= hoy):
            resultado.append(fila)
    return resultado


def _materias(id_carrera=None):
    conexion = conectar()
    try:
        sql = ("SELECT m.id_materia, m.nom_materia, m.cant_horas, c.id_carrera, c.nom_carrera, "
               "c.resolucion FROM materia m JOIN carrera c ON c.id_carrera = m.id_carrera")
        parametros = ()
        if id_carrera:
            sql += " WHERE c.id_carrera = ?"
            parametros = (id_carrera,)
        return conexion.execute(sql + " ORDER BY c.nom_carrera, m.nom_materia", parametros).fetchall()
    finally:
        conexion.close()


def planta_docente(id_carrera=None, hoy=None):
    """[{carrera, resolucion, materias: [{materia, horas, docentes: [...]}]}] con la situación
    de hoy. Cada docente trae 'estado' (texto) y 'ausente' (bool)."""
    hoy = hoy or date.today()
    vigentes = _designaciones_vigentes(hoy)
    ausentes = ausentes_hoy(hoy)

    carreras = {}
    for id_materia, materia, horas, id_carrera_, carrera, resolucion in _materias(id_carrera):
        bloque = carreras.setdefault(id_carrera_, {"carrera": carrera, "resolucion": resolucion,
                                                   "materias": []})
        docentes = []
        for v in (x for x in vigentes if x["id_materia"] == id_materia):
            ausencia = ausentes.get(v["id_docente"])
            if ausencia:
                estado = f"{ausencia[0]} hasta {ausencia[1].strftime(FORMATO_FECHA)}"
            elif v["reemplaza"]:
                estado = f"Suplente de {v['reemplaza']}"
            else:
                estado = "En funciones"
            docentes.append(dict(v, estado=estado, ausente=bool(ausencia)))
        docentes.sort(key=lambda x: (x["situacion"] == "S", x["apellido"]))
        bloque["materias"].append({"materia": materia, "horas": horas, "docentes": docentes})
    return list(carreras.values())


def materias_sin_cubrir(hoy=None):
    """(vacantes, descubiertas):
    - vacantes: materias sin ninguna designación vigente.
    - descubiertas: materias cuyos docentes están todos de licencia/carpeta médica hoy.
    Cada elemento: {carrera, materia, horas, detalle}."""
    vacantes, descubiertas = [], []
    for bloque in planta_docente(hoy=hoy):
        for m in bloque["materias"]:
            base = {"carrera": bloque["carrera"], "materia": m["materia"], "horas": m["horas"]}
            if not m["docentes"]:
                vacantes.append(dict(base, detalle="Sin docente designado"))
            elif all(d["ausente"] for d in m["docentes"]):
                detalle = "; ".join(f"{d['apellido']}, {d['nombre']} ({d['estado']})"
                                    for d in m["docentes"])
                descubiertas.append(dict(base, detalle=detalle))
    return vacantes, descubiertas


# ---------- PDF ----------

_CELDA = ParagraphStyle("planta_celda", fontName="Helvetica", fontSize=8, leading=10)
_CELDA_C = ParagraphStyle("planta_celda_c", parent=_CELDA, alignment=TA_CENTER)
_CABECERA = ParagraphStyle("planta_cab", fontName="Helvetica-Bold", fontSize=8, leading=10,
                           alignment=TA_CENTER, textColor=colors.white)
_CARRERA = ParagraphStyle("planta_carrera", fontName="Helvetica-Bold", fontSize=11, leading=14)
_RESUMEN = ParagraphStyle("planta_resumen", fontName="Helvetica", fontSize=9, leading=12,
                          textColor=colors.HexColor("#444444"))
_TITULO = ParagraphStyle("planta_tit", fontName="Helvetica-Bold", fontSize=14, leading=18,
                         alignment=TA_CENTER)

COLUMNAS = ["Materia", "Hs", "Año", "Div.", "Docente", "DNI", "S", "C", "T",
            "Toma de posesión", "Baja", "S.D.N°", "Situación actual"]
ANCHOS = [4.6, 0.8, 0.9, 0.9, 4.0, 1.8, 0.6, 0.8, 0.6, 1.9, 1.9, 1.5, 4.9]


def _texto(valor):
    from html import escape
    return escape(str(valor)) if valor not in (None, "") else "-"


def _filas_carrera(bloque):
    filas, estilos = [], []
    for m in bloque["materias"]:
        if not m["docentes"]:
            filas.append([Paragraph(_texto(m["materia"]), _CELDA), Paragraph(_texto(m["horas"]), _CELDA_C)]
                         + [""] * 2 + [Paragraph("<font color='#C62828'><b>VACANTE</b></font>", _CELDA)]
                         + [""] * 8)
            estilos.append(("BACKGROUND", (0, len(filas)), (-1, len(filas)), colors.HexColor("#FDECEA")))
            continue
        for i, d in enumerate(m["docentes"]):
            estado = _texto(d["estado"])
            if d["ausente"]:
                estado = f"<font color='#C62828'>{estado}</font>"
            filas.append([
                Paragraph(_texto(m["materia"]) if i == 0 else "", _CELDA),
                Paragraph(_texto(m["horas"]) if i == 0 else "", _CELDA_C),
                Paragraph(_texto(d["anio"]), _CELDA_C), Paragraph(_texto(d["division"]), _CELDA_C),
                Paragraph(_texto(f"{str(d['apellido']).upper()}, {d['nombre']}"), _CELDA),
                Paragraph(_texto(d["dni"]), _CELDA_C), Paragraph(_texto(d["situacion"]), _CELDA_C),
                Paragraph(_texto(d["condicion"]), _CELDA_C), Paragraph(_texto(d["turno"]), _CELDA_C),
                Paragraph(_texto(d["alta"]), _CELDA_C), Paragraph(_texto(d["baja"]), _CELDA_C),
                Paragraph(_texto(d["disposicion"]), _CELDA_C), Paragraph(estado, _CELDA),
            ])
    return filas, estilos


def exportar_planta_pdf(ruta, id_carrera=None):
    from certificados.comun import encabezado_institucional, fecha_en_letras

    hoy = date.today()
    margen = 1.2 * cm
    ancho_util = landscape(A4)[0] - 2 * margen
    escala = ancho_util / (sum(ANCHOS) * cm)
    contenido = encabezado_institucional(ancho_util) + [
        Paragraph("PLANTA DOCENTE", _TITULO),
        Paragraph(f"Situación al {fecha_en_letras(hoy)}", ParagraphStyle(
            "planta_fecha", parent=_RESUMEN, alignment=TA_CENTER)),
        Spacer(1, 0.4 * cm),
    ]
    bloques = planta_docente(id_carrera, hoy)
    if not bloques:
        contenido.append(Paragraph("No hay materias cargadas.", _RESUMEN))

    for bloque in bloques:
        materias = bloque["materias"]
        cubiertas = sum(1 for m in materias if m["docentes"])
        horas = sum(int(m["horas"] or 0) for m in materias)
        resolucion = str(bloque["resolucion"] or "").strip()
        if resolucion and not resolucion.lower().startswith("res"):
            resolucion = f"Res. Min. {resolucion}"
        titulo = _texto(bloque["carrera"]) + (f" — {_texto(resolucion)}" if resolucion else "")
        filas, estilos = _filas_carrera(bloque)
        tabla = Table([[Paragraph(c, _CABECERA) for c in COLUMNAS]] + filas,
                      colWidths=[a * cm * escala for a in ANCHOS], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#888888")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ] + estilos))
        contenido += [
            Paragraph(titulo, _CARRERA),
            Paragraph(f"{len(materias)} materia(s) · {cubiertas} cubierta(s) · "
                      f"{len(materias) - cubiertas} vacante(s) · {horas} horas cátedra en total",
                      _RESUMEN),
            Spacer(1, 0.15 * cm), tabla, Spacer(1, 0.6 * cm),
        ]
    contenido.append(Paragraph("S: situación (I = interino, S = suplente) · C: condición "
                               "(1C/2C = cuatrimestre, A = anual) · T: turno", _RESUMEN))
    SimpleDocTemplate(ruta, pagesize=landscape(A4), leftMargin=margen, rightMargin=margen,
                      topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                      title="Planta docente").build(contenido)


# ---------- EXCEL ----------

def exportar_planta_excel(ruta, id_carrera=None):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Planta docente"
    encabezados = ["Carrera", "Resolución", "Materia", "Horas", "Año", "División", "Docente",
                   "DNI", "Situación", "Condición", "Turno", "Toma de posesión", "Baja",
                   "S.D.N°", "Situación actual"]
    hoja.append(encabezados)
    rojo = PatternFill("solid", fgColor="FDECEA")
    for bloque in planta_docente(id_carrera):
        for m in bloque["materias"]:
            base = [bloque["carrera"], bloque["resolucion"], m["materia"], m["horas"]]
            if not m["docentes"]:
                hoja.append(base + ["", "", "VACANTE"] + [""] * 8)
                for celda in hoja[hoja.max_row]:
                    celda.fill = rojo
                continue
            for d in m["docentes"]:
                hoja.append(base + [d["anio"], d["division"], f"{d['apellido']}, {d['nombre']}",
                                    d["dni"], d["situacion"], d["condicion"], d["turno"],
                                    d["alta"], d["baja"], d["disposicion"], d["estado"]])
                if d["ausente"]:
                    for celda in hoja[hoja.max_row]:
                        celda.fill = rojo

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1565C0")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for i, columna in enumerate(hoja.columns, start=1):
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 3, 8), 45)
    libro.save(ruta)
