"""Exportación del listado de docentes a Excel."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database.conexion import conectar
from database.consultas import listar_docentes


def _formatear_hoja(hoja):
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1565C0")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for i, columna in enumerate(hoja.columns, start=1):
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 3, 10), 50)


def exportar_docentes_excel(ruta):
    libro = Workbook()

    hoja = libro.active
    hoja.title = "Docentes"
    hoja.append(["DNI", "Apellido", "Nombre", "CUIL", "Fecha de nacimiento",
                 "Domicilio", "Email", "Carácter", "Estado"])
    for d in listar_docentes("", incluir_inactivos=True):
        hoja.append([d[0], d[2], d[1], d[3], d[4], d[5], d[6], d[7],
                     "Activo" if d[8] else "Inactivo"])
    _formatear_hoja(hoja)

    hoja2 = libro.create_sheet("Designaciones")
    hoja2.append(["DNI", "Apellido", "Nombre", "Carrera", "Materia", "Disposición",
                  "Fecha de alta", "Fecha de baja", "Motivo de baja"])
    conexion = conectar()
    try:
        filas = conexion.execute(
            "SELECT doc.dni, doc.apellido, doc.nombre, c.nom_carrera, m.nom_materia, "
            "de.nro_disposicion, de.fecha_alta, de.fecha_baja, de.motivo_baja "
            "FROM designacion de "
            "JOIN docentes doc ON doc.id_docente = de.id_docente "
            "JOIN materia m ON m.id_materia = de.id_materia "
            "JOIN carrera c ON c.id_carrera = m.id_carrera "
            "ORDER BY doc.apellido, doc.nombre, c.nom_carrera, m.nom_materia"
        ).fetchall()
    finally:
        conexion.close()
    for fila in filas:
        hoja2.append(list(fila))
    _formatear_hoja(hoja2)

    libro.save(ruta)
