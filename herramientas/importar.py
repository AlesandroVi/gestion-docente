"""Importar docentes desde una planilla de Excel (y generar la planilla modelo)."""
import sqlite3
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import CARACTERES
from crud.validaciones import NOMBRES_CAMPOS, campo_duplicado, validar_docente
from database.conexion import conectar
from database.consultas import FORMATO_FECHA

COLUMNAS = ["Apellido", "Nombre", "DNI", "CUIL", "Fecha de nacimiento (dd/mm/aaaa)",
            "Domicilio", "Email", "Carácter"]
CLAVES = ["apellido", "nombre", "dni", "cuil", "fecha_nacimiento", "domicilio", "email",
          "caracter"]


def crear_plantilla(ruta):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Docentes"
    hoja.append(COLUMNAS)
    hoja.append(["Pérez", "Juan Carlos", 30123456, 20301234565, "01/01/1990",
                 "Calle Falsa 123", "juan@correo.com", CARACTERES[0]])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1565C0")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for celda in hoja[2]:
        celda.font = Font(italic=True, color="888888")
    for letra, ancho in zip("ABCDEFGH", (22, 22, 12, 15, 20, 30, 28, 14)):
        hoja.column_dimensions[letra].width = ancho
    hoja.row_dimensions[1].height = 32
    hoja.freeze_panes = "A2"
    # Lista desplegable para el carácter
    lista = DataValidation(type="list", formula1='"' + ",".join(CARACTERES) + '"', allow_blank=True)
    hoja.add_data_validation(lista)
    lista.add("H2:H1000")

    ayuda = libro.create_sheet("Instrucciones")
    for linea in [
        "Cómo completar la planilla:",
        "• Un docente por fila, empezando en la fila 2 (la fila gris de ejemplo se puede borrar o pisar).",
        "• DNI y CUIL: solo números, sin puntos ni guiones.",
        "• Fecha de nacimiento: dd/mm/aaaa (ej: 12/05/1994) o una celda con formato de fecha.",
        f"• Carácter: {' o '.join(CARACTERES)}.",
        "• Los docentes que ya estén cargados (mismo DNI, CUIL o email) se saltean.",
        "• Las filas con errores no se cargan: el sistema te dice qué corregir en cada una.",
    ]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=12)
    ayuda.column_dimensions["A"].width = 100
    libro.save(ruta)


def _celda_a_texto(valor, clave):
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime(FORMATO_FECHA)
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if clave in ("dni", "cuil"):
        texto = texto.replace(".", "").replace("-", "").replace(" ", "")
    if clave == "caracter":
        for opcion in CARACTERES:
            if texto.lower() == opcion.lower():
                return opcion
    return texto


def importar_docentes(ruta):
    """Devuelve (cantidad importada, [(fila, nombre, mensaje)] con lo que no se cargó)."""
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hoja = libro["Docentes"] if "Docentes" in libro.sheetnames else libro.worksheets[0]

    importados, problemas = 0, []
    conexion = conectar()
    try:
        for numero, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            valores = list(fila[:len(CLAVES)]) + [None] * (len(CLAVES) - len(fila[:len(CLAVES)]))
            if all(v in (None, "") for v in valores):
                continue
            datos = {c: _celda_a_texto(v, c) for c, v in zip(CLAVES, valores)}
            if datos["dni"] == "30123456" and datos["email"] == "juan@correo.com":
                continue  # fila de ejemplo de la plantilla
            nombre = f"{datos['apellido']}, {datos['nombre']}".strip(", ")
            errores, limpios = validar_docente(
                datos["nombre"], datos["apellido"], datos["dni"], datos["cuil"],
                datos["fecha_nacimiento"], datos["domicilio"], datos["email"], datos["caracter"],
                caracteres_validos=CARACTERES)
            if errores:
                detalle = "; ".join(f"{NOMBRES_CAMPOS.get(k, k)}: {m}" for k, m in errores.items())
                problemas.append((numero, nombre, detalle))
                continue
            try:
                with conexion:
                    conexion.execute(
                        "INSERT INTO docentes (nombre, apellido, dni, cuil, fecha_nacimiento, "
                        "domicilio, email, caracter) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                        (limpios["nombre"], limpios["apellido"], limpios["dni"], limpios["cuil"],
                         limpios["fecha_nacimiento"], limpios["domicilio"], limpios["email"],
                         limpios["caracter"]))
                importados += 1
            except sqlite3.IntegrityError as error:
                duplicado = campo_duplicado(error)
                problemas.append((numero, nombre, duplicado[1] if duplicado else str(error)))
    finally:
        conexion.close()
        libro.close()
    return importados, problemas
